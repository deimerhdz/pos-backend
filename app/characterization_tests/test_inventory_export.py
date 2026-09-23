"""Tests de la nueva funcionalidad — spec 086-exportacion-excel-inventario:
GET /inventory/items/export (Historia 1: descarga del reporte completo;
Historia 2: consistencia de columnas y datos exportados).

No son characterization tests -- el endpoint no existe antes de este spec
(comportamiento enteramente nuevo, ver `research.md`/`contracts/
inventory-export-api.md` del propio spec). Combina las tablas de
`plan_fixtures.py` (tenants/plans/roles/users, para el gating de plan y el
rol ADMIN) con las de `fixtures.py` (insumos/unidades de medida), mismo
patrón que `test_plan_gating_inventory_fields.py`.

Ejecutar solo este módulo:

    python -m unittest app.characterization_tests.test_inventory_export -v
"""
from __future__ import annotations

import io
import unittest
from decimal import Decimal
from unittest import mock

import openpyxl
from fastapi import FastAPI
from sqlalchemy import create_engine
from sqlalchemy.orm import Session
from starlette.testclient import TestClient

from app.characterization_tests import fixtures as f
from app.characterization_tests import plan_fixtures as pf
from app.api.v1.inventory.router import router as inventory_router
from app.core.db import get_db, get_tenant
from app.core.dependencies import get_current_user
from app.core.models import Base

_TABLE_NAMES = [
    # plan_fixtures.py — gating de plan y rol ADMIN
    "tenants", "plans", "roles", "users",
    # fixtures.py — insumos y unidades de medida
    "inventory_items", "inventory_movements", "unit_measures",
]


def _new_session() -> Session:
    """`check_same_thread=False`: a diferencia de las demás fixtures del
    paquete, este módulo ejercita el endpoint vía `TestClient` (spec 086),
    que ejecuta el handler en un thread distinto al del test — mismo ajuste
    que `test_notifications.py` (spec 077) ya necesitó por el mismo motivo."""
    tables = [t for t in Base.metadata.tables.values() if t.name in _TABLE_NAMES]
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    conn = engine.connect().execution_options(
        schema_translate_map={"tenant": None, "shared": None}
    )
    Base.metadata.create_all(bind=conn, tables=tables)
    conn.commit()
    return Session(bind=conn)


def _build_client(db: Session, tenant, user) -> TestClient:
    """App mínima con el router real de inventario -- mismo patrón que
    `test_notifications.py` (spec 077): se sobreescriben `get_db`,
    `get_current_user` y `get_tenant` (esta última la usa
    `require_module_access("inventario")`, aplicado a nivel de router)."""
    app = FastAPI()
    app.include_router(inventory_router, prefix="/api/v1")
    app.dependency_overrides[get_db] = lambda: db
    app.dependency_overrides[get_current_user] = lambda: user
    app.dependency_overrides[get_tenant] = lambda: tenant
    return TestClient(app, raise_server_exceptions=False)


def _con_modulo_inventario(db: Session):
    plan = pf.make_plan(db, inventario_access=True)
    tenant = pf.make_tenant(db, plan=plan)
    return tenant


class ExportEndpointTests(unittest.TestCase):
    """Historia 1 (P1): descarga del reporte completo."""

    def setUp(self):
        self.db = _new_session()
        self.tenant = _con_modulo_inventario(self.db)

    def test_admin_con_insumos_descarga_n_mas_1_filas(self):
        for _ in range(3):
            f.make_inventory_item(self.db)
        f.make_inventory_item(self.db, active=False)
        self.db.commit()

        admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")
        self.db.commit()
        client = _build_client(self.db, self.tenant, admin)

        resp = client.get("/api/v1/inventory/items/export")

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn('filename="inventario_', resp.headers["content-disposition"])
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.active.max_row, 4 + 1)  # 4 insumos (3 activos + 1 inactivo)

    def test_inventario_vacio_descarga_solo_el_encabezado(self):
        admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")
        self.db.commit()
        client = _build_client(self.db, self.tenant, admin)

        resp = client.get("/api/v1/inventory/items/export")

        self.assertEqual(resp.status_code, 200)
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertEqual(wb.active.max_row, 1)

    def test_usuario_sin_rol_admin_recibe_403_sin_archivo(self):
        cashier = pf.make_user(self.db, self.tenant, role_name="CASHIER")
        self.db.commit()
        client = _build_client(self.db, self.tenant, cashier)

        resp = client.get("/api/v1/inventory/items/export")

        self.assertEqual(resp.status_code, 403)
        self.assertNotIn("content-disposition", resp.headers)

    def test_fallo_a_mitad_de_generacion_responde_500_sin_archivo_parcial(self):
        admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")
        self.db.commit()
        client = _build_client(self.db, self.tenant, admin)

        with mock.patch(
            "app.api.v1.inventory.router.build_inventory_excel",
            side_effect=RuntimeError("fallo simulado a mitad de la generación"),
        ):
            resp = client.get("/api/v1/inventory/items/export")

        self.assertEqual(resp.status_code, 500)
        self.assertNotIn("content-disposition", resp.headers)
        self.assertNotIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            resp.headers.get("content-type", ""),
        )


class ExportColumnsTests(unittest.TestCase):
    """Historia 2 (P2): encabezados exactos y celdas numéricas nativas."""

    def setUp(self):
        self.db = _new_session()
        self.tenant = _con_modulo_inventario(self.db)
        self.admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")

    def _export(self):
        self.db.commit()
        client = _build_client(self.db, self.tenant, self.admin)
        resp = client.get("/api/v1/inventory/items/export")
        self.assertEqual(resp.status_code, 200)
        return openpyxl.load_workbook(io.BytesIO(resp.content))

    def test_fila_1_son_los_encabezados_exactos_en_orden(self):
        wb = self._export()
        headers = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(headers, ["Nombre", "Tipo", "Unidad", "Stock", "Mínimo", "Costo", "Estado"])

    def test_celdas_stock_minimo_costo_son_numericas_no_texto(self):
        f.make_inventory_item(
            self.db, current_stock=Decimal("12.500"), min_stock=Decimal("3.000"),
            unit_cost=Decimal("15.00"),
        )
        wb = self._export()
        row = next(wb.active.iter_rows(min_row=2, max_row=2))
        stock_cell, min_cell, cost_cell = row[3], row[4], row[5]
        self.assertEqual(stock_cell.data_type, "n")
        self.assertEqual(min_cell.data_type, "n")
        self.assertEqual(cost_cell.data_type, "n")

    def test_costo_15_exporta_numero_15_no_texto_con_simbolo_de_moneda(self):
        f.make_inventory_item(self.db, unit_cost=Decimal("15.00"))
        wb = self._export()
        row = next(wb.active.iter_rows(min_row=2, max_row=2))
        cost_cell = row[5]
        self.assertEqual(float(cost_cell.value), 15.0)
        self.assertNotIsInstance(cost_cell.value, str)


class ExportLabelsTests(unittest.TestCase):
    """Historia 2 (P2): etiquetas de Tipo/Unidad/Estado idénticas a pantalla."""

    def setUp(self):
        self.db = _new_session()
        self.tenant = _con_modulo_inventario(self.db)
        self.admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")

    def _export(self):
        self.db.commit()
        client = _build_client(self.db, self.tenant, self.admin)
        resp = client.get("/api/v1/inventory/items/export")
        return openpyxl.load_workbook(io.BytesIO(resp.content))

    def test_tipo_raw_material_exporta_materia_prima_y_packaged_exporta_empacado(self):
        f.make_inventory_item(self.db, name="Leche", type="raw_material")
        f.make_inventory_item(self.db, name="Vaso desechable", type="packaged")
        wb = self._export()
        tipos = {row[0].value: row[1].value for row in wb.active.iter_rows(min_row=2)}
        self.assertEqual(tipos["Leche"], "Materia prima")
        self.assertEqual(tipos["Vaso desechable"], "Empacado")

    def test_unidad_exporta_la_abreviatura_no_el_nombre_largo(self):
        unit = f.make_unit(self.db, name="Litro", abbreviation="lt")
        f.make_inventory_item(self.db, name="Leche", unit=unit)
        wb = self._export()
        row = next(wb.active.iter_rows(min_row=2, max_row=2))
        self.assertEqual(row[2].value, "lt")

    def test_activo_bajo_minimo_exporta_bajo_minimo_inactivo_igual_exporta_ok(self):
        f.make_inventory_item(
            self.db, name="Activo bajo mínimo",
            active=True, current_stock=Decimal("1"), min_stock=Decimal("5"),
        )
        f.make_inventory_item(
            self.db, name="Inactivo bajo mínimo",
            active=False, current_stock=Decimal("1"), min_stock=Decimal("5"),
        )
        wb = self._export()
        estados = {row[0].value: row[6].value for row in wb.active.iter_rows(min_row=2)}
        self.assertEqual(estados["Activo bajo mínimo"], "Bajo mínimo")
        self.assertEqual(estados["Inactivo bajo mínimo"], "OK")


class ExportEncodingTests(unittest.TestCase):
    """Historia 2 (P2): nombres con tildes/eñes/comillas sin corrupción (FR-007)."""

    def setUp(self):
        self.db = _new_session()
        self.tenant = _con_modulo_inventario(self.db)
        self.admin = pf.make_user(self.db, self.tenant, role_name="ADMIN")

    def test_nombres_con_caracteres_del_espanol_se_preservan_intactos(self):
        f.make_inventory_item(self.db, name="Champiñón")
        f.make_inventory_item(self.db, name="Arequipe (500ml)")
        self.db.commit()
        client = _build_client(self.db, self.tenant, self.admin)
        resp = client.get("/api/v1/inventory/items/export")

        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        nombres = [row[0].value for row in wb.active.iter_rows(min_row=2)]
        self.assertIn("Champiñón", nombres)
        self.assertIn("Arequipe (500ml)", nombres)


if __name__ == "__main__":
    unittest.main()
